# -*- coding: utf-8 -*-
"""比较 Excel 中相邻记录的 speed，并标记超过阈值的两端单元格。"""
import os
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, font

from openpyxl import load_workbook
from openpyxl.styles import Font


def find_speed_column(ws):
    """在前 20 行中定位 speed 表头，支持大小写和首尾空格差异。"""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip().lower() == "speed":
                return cell.column, cell.row
    return None, None


def mark_speed_changes(input_path, threshold):
    wb = load_workbook(input_path)
    matched = 0
    compared = 0

    for ws in wb.worksheets:
        col, header_row = find_speed_column(ws)
        if col is None:
            continue

        previous = None
        for row_idx in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col)
            value = cell.value
            if value is None or (isinstance(value, str) and not value.strip()):
                previous = None
                continue
            try:
                speed = float(value)
            except (TypeError, ValueError):
                previous = None
                continue

            if previous is not None:
                compared += 1
                prev_cell, prev_speed = previous
                if abs(speed - prev_speed) > threshold:
                    for target in (prev_cell, cell):
                        target.font = target.font.copy(color="FFFF0000", bold=True)
                    matched += 1
            previous = (cell, speed)

    if matched == 0 and compared == 0:
        raise ValueError("没有找到可比较的 speed 数据列。请确认表头为 speed 且下面有数字数据。")

    base, ext = os.path.splitext(input_path)
    output_path = f"{base}_speed对比{ext}"
    wb.save(output_path)
    return output_path, compared, matched


def main():
    root = tk.Tk()
    root.withdraw()
    input_path = filedialog.askopenfilename(
        title="选择需要对比的 Excel 文件",
        filetypes=[("Excel 文件", "*.xlsx *.xlsm"), ("所有文件", "*.*")],
    )
    if not input_path:
        root.destroy()
        return

    threshold = simpledialog.askfloat(
        "设置速度差范围",
        "请输入相邻两行 speed 的差值范围：\n差值绝对值大于此数时，两格都会标红。",
        minvalue=0,
        parent=root,
    )
    if threshold is None:
        root.destroy()
        return

    try:
        output_path, compared, matched = mark_speed_changes(input_path, threshold)
        messagebox.showinfo(
            "对比完成",
            f"已比较 {compared} 组相邻 speed 数据。\n"
            f"超过范围：{matched} 组。\n\n结果文件：\n{output_path}",
            parent=root,
        )
    except Exception as exc:
        messagebox.showerror("无法处理文件", str(exc), parent=root)
    finally:
        root.destroy()


if __name__ == "__main__":
    main()