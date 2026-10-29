# -*- coding: utf-8 -*-
"""批量将 OziExplorer .plt 轨迹文件转换为 Excel(.xlsx)。"""
import datetime
import os
import sys
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def choose_files():
    """选择一个或多个 PLT 文件。"""
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        paths = filedialog.askopenfilenames(
            title="选择要转换的 PLT 文件",
            filetypes=[("PLT 文件", "*.plt"), ("所有文件", "*.*")],
        )
        root.destroy()
        return list(paths)
    except Exception:
        return []


def oadate_to_str(serial):
    try:
        return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(serial))).strftime("%Y-%m-%d %H:%M:%S")
    except (ValueError, TypeError, OverflowError):
        return ""


def convert_one(src_path, out_path):
    with open(src_path, "r", encoding="utf-8-sig") as f:
        lines = f.read().splitlines()
    if len(lines) <= 6:
        raise ValueError("文件格式错误")

    data = []
    for line in lines[6:]:
        if not line.strip():
            continue
        parts = [part.strip() for part in line.split(",")]
        if len(parts) < 5:
            raise ValueError("文件格式错误")
        try:
            lat, lon, alt_ft = float(parts[0]), float(parts[1]), float(parts[3])
            float(parts[4])
        except ValueError:
            raise ValueError("文件格式错误") from None
        data.append((lat, lon, parts[2], alt_ft, parts[4]))
    if not data:
        raise ValueError("文件格式错误")

    wb = Workbook()
    ws = wb.active
    ws.title = "轨迹点"
    ws.append(["序号", "纬度", "经度", "标志位", "高程(英尺)", "高程(米)", "时间(OADate)", "时间(本地)"])
    for i, (lat, lon, flag, alt_ft, timestamp) in enumerate(data, 1):
        ws.append([i, lat, lon, flag, alt_ft, round(alt_ft * 0.3048, 2), timestamp, oadate_to_str(timestamp)])
    for idx, width in enumerate([8, 14, 14, 8, 12, 12, 18, 20], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    wb.save(out_path)
    times = [oadate_to_str(row[4]) for row in data]
    times = [value for value in times if value]
    return len(data), times[0] if times else "", times[-1] if times else ""


def main():
    args = [arg.strip() for arg in sys.argv[1:] if arg.strip()]
    if not args:
        files = choose_files()
        if not files:
            print("未选择文件，已退出。")
            return
        out_dir = os.path.dirname(os.path.abspath(files[0]))
    elif os.path.isfile(args[0]):
        files = [args[0]]
        out_dir = args[1] if len(args) > 1 else os.path.dirname(os.path.abspath(files[0]))
    elif os.path.isdir(args[0]):
        print("请直接选择 PLT 文件，不再搜索文件夹。")
        return
    else:
        print(f"找不到文件：{args[0]}")
        return

    os.makedirs(out_dir, exist_ok=True)
    for src in files:
        out = os.path.join(out_dir, os.path.splitext(os.path.basename(src))[0] + ".xlsx")
        try:
            count, start, end = convert_one(src, out)
            print(f"[成功] {os.path.basename(src)} -> {out}，{count} 点，时间 {start} ~ {end}")
        except (OSError, UnicodeError, ValueError) as ex:
            print(f"[文件格式错误] {src}: {ex}")


if __name__ == "__main__":
    main()