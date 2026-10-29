#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
GPS 数据「延时 / 丢失」分析工具
================================
功能：
  1) 新增「延时」列：逐行计算 接收时间 - 卫星时间（单位秒），
     每行都填入，格式 ±H:MM:SS（例 0:01:50 / -0:00:05）；
     接收时间早于卫星时间时显示负值，输出 Excel 中负值标红加粗。
  2) 新增「丢失」列：计算「卫星时间」相邻行的间隔（秒），
     间隔 > 你输入的阈值（秒）时，标记为「丢失」，否则「正常」。

用法：
  A. 直接双击/无参数运行 -> 弹窗选文件 + 弹窗输入阈值
  B. 命令行：
     python analyze_gps_delay.py --input 数据.xlsx --threshold 4
     （threshold 省略则弹窗输入，默认 4 秒）

依赖：pandas, openpyxl
安装：pip install pandas openpyxl
"""
import os
import sys
import argparse
import pandas as pd

# ---------- 列名配置（如你的文件列名不同，改这里即可） ----------
SAT_COL_CANDIDATES = ["卫星时间", "卫星", "satellite", "gps_time"]
RECV_COL_CANDIDATES = ["接收时间", "接收", "recv", "receive_time"]
# 延时列不再做阈值过滤：每行都计算（接收 - 卫星），负值表示接收早于卫星
# -----------------------------------------------------------------

def find_col(df, *candidates):
    for c in df.columns:
        cs = str(c)
        for kw in candidates:
            if kw.lower() in cs.lower():
                return c
    return None

def fmt_hms(seconds):
    """把秒数格式化为 ±H:MM:SS（保留正负号，用于负延时）"""
    seconds = int(round(seconds))
    sign = "-" if seconds < 0 else ""
    seconds = abs(seconds)
    h = seconds // 3600
    m = (seconds % 3600) // 60
    s = seconds % 60
    return f"{sign}{h}:{m:02d}:{s:02d}"

def pick_input_file():
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title="选择 GPS 数据文件",
        filetypes=[("Excel/CSV", "*.xlsx;*.xls;*.csv"), ("All", "*.*")])
    root.destroy()
    return path

def ask_threshold(default=4):
    import tkinter as tk
    from tkinter import simpledialog
    root = tk.Tk()
    root.withdraw()
    val = simpledialog.askfloat(
        "相邻间隔阈值",
        "请输入「卫星时间」相邻行允许的最大间隔（秒）：\n超过该值将标记为「丢失」",
        initialvalue=default)
    root.destroy()
    return val

def load_df(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xls"):
        return pd.read_excel(path)
    return pd.read_csv(path)

def mark_negative_delay(path):
    """把输出 Excel「延时」列中的负值标红加粗 + 浅红底色（仅 xlsx）"""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    wb = load_workbook(path)
    ws = wb.active
    # 找到表头中的「延时」列
    delay_col = None
    for cell in ws[1]:
        if str(cell.value) == "延时":
            delay_col = cell.column
            break
    if delay_col is None:
        return
    red_font = Font(color="FF0000", bold=True)
    red_fill = PatternFill(fill_type="solid", start_color="FFFFC7CE", end_color="FFFFC7CE")
    n_neg = 0
    for row in range(2, ws.max_row + 1):
        c = ws.cell(row=row, column=delay_col)
        if isinstance(c.value, str) and c.value.startswith("-"):
            c.font = red_font
            c.fill = red_fill
            n_neg += 1
    wb.save(path)
    return n_neg

def analyze(df, threshold, sort_by_time=False):
    sat_col = find_col(df, *SAT_COL_CANDIDATES)
    recv_col = find_col(df, *RECV_COL_CANDIDATES)
    if not sat_col or not recv_col:
        raise RuntimeError(
            f"未找到时间列。现有列：{list(df.columns)}\n"
            f"请确认文件含「卫星时间」「接收时间」列，或在脚本顶部修改 SAT_COL_CANDIDATES / RECV_COL_CANDIDATES。")

    df[sat_col] = pd.to_datetime(df[sat_col], errors="coerce")
    df[recv_col] = pd.to_datetime(df[recv_col], errors="coerce")

    if sort_by_time:
        df = df.sort_values(sat_col, kind="stable").reset_index(drop=True)

    # 1) 延时：每行都计算 接收 - 卫星，保留正负号（接收早于卫星即为负）
    delta = (df[recv_col] - df[sat_col]).dt.total_seconds()
    df["延时"] = ""
    valid = delta.notna()
    df.loc[valid, "延时"] = delta[valid].apply(fmt_hms)

    # 2) 丢失：相邻卫星时间间隔 > 阈值
    gap = df[sat_col].diff().dt.total_seconds()   # 首行为 NaN
    df["丢失"] = ""
    gap_mask = gap.abs() > threshold
    df.loc[gap_mask, "丢失"] = "丢失"
    df.loc[~gap_mask & gap.notna(), "丢失"] = "正常"
    return df

def main():
    parser = argparse.ArgumentParser(description="GPS 延时/丢失分析")
    parser.add_argument("--input", help="输入 xlsx/csv 路径")
    parser.add_argument("--threshold", type=float, help="相邻间隔阈值(秒)")
    parser.add_argument("--output", help="输出路径(可选)")
    args = parser.parse_args()

    input_path = args.input
    if not input_path:
        try:
            input_path = pick_input_file()
        except Exception:
            input_path = input("请输入数据文件路径：").strip('"')
    if not input_path or not os.path.exists(input_path):
        print("未选择有效文件，退出。")
        sys.exit(1)

    threshold = args.threshold
    if threshold is None:
        try:
            t = ask_threshold(4)
            threshold = t if t is not None else 4.0
        except Exception:
            threshold = float(input("请输入相邻间隔阈值(秒)，默认 4：") or 4)

    df = load_df(input_path)
    df = analyze(df, threshold)

    out = args.output
    if not out:
        base, ext = os.path.splitext(input_path)
        out = base + "_分析.xlsx"
    n_neg = 0
    if out.lower().endswith((".xlsx", ".xls")):
        df.to_excel(out, index=False)
        if out.lower().endswith(".xlsx"):
            n_neg = mark_negative_delay(out) or 0
    else:
        df.to_csv(out, index=False, encoding="utf-8-sig")  # CSV 无法保存标红样式

    n_delay = int((df["延时"] != "").sum())
    n_lost = int((df["丢失"] == "丢失").sum())
    print(f"处理完成：共 {len(df)} 行")
    print(f"  延时(已计算) 行数：{n_delay}")
    print(f"  延时为负(标红) 行数：{n_neg}")
    print(f"  丢失(间隔> {threshold}s) 行数：{n_lost}")
    print(f"  输出文件：{out}")

if __name__ == "__main__":
    main()
