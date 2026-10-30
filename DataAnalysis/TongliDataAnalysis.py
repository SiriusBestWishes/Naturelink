#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
GPS 数据「延时 / 丢失」分析工具
================================
功能：
  1) 新增「延时」列：逐行计算 接收时间 - 卫星时间（单位秒），
     每行都填入，格式 ±H:MM:SS（例 0:01:50 / -0:00:05）；
     接收时间早于卫星时间时显示负值，输出 Excel 中负值标红加粗。
  2) 新增「丢失」列：先读「ACC」车辆状态列，按状态分档判断相邻上报间隔——
     ACC 开超过 阈值A（弹窗输入，默认 15 秒）、ACC 关超过 阈值B（弹窗输入，默认 120 秒）
     时标记为「丢失」（标黄加粗），否则「正常」。
  3) 新增「重发」列：逐行比较「卫星时间」，若与上一行相同则标记为「重发」。
  4) 输出 Excel 自动适配列宽（中文按双宽计算），并冻结首行表头。
  5) 控制台统计：总行数、延时为负、丢失（按 ACC 分档）、重发、定位/未定位 各自条数与占比。

用法：
  A. 直接双击/无参数运行 -> 弹窗选文件 + 弹窗输入 ACC 开/关两个阈值
  B. 命令行：
     python TongliDataAnalysis.py --input 数据.xlsx --acc-on 15 --acc-off 120
     （--acc-on / --acc-off 省略则弹窗输入，默认 15 / 120 秒）

依赖：pandas, openpyxl, numpy
安装：pip install pandas openpyxl numpy
"""
import os
import sys
import time
import argparse
import numpy as np
import pandas as pd

# ---------- 列名配置（如你的文件列名不同，改这里即可） ----------
SAT_COL_CANDIDATES = ["卫星时间", "卫星", "satellite", "gps_time"]
RECV_COL_CANDIDATES = ["接收时间", "接收", "recv", "receive_time"]
ACC_COL_CANDIDATES = ["ACC", "车辆状态", "acc状态", "acc"]
POS_COL_CANDIDATES = ["GNSS信号", "GNSS", "定位状态", "定位", "是否定位", "gps状态", "position", "fix"]
# 延时列不再做阈值过滤：每行都计算（接收 - 卫星），负值表示接收早于卫星
# -----------------------------------------------------------------

def find_col(df, *candidates):
    # 1) 先精确匹配，避免「定位」误命中「定位时间」这类列
    for kw in candidates:
        for c in df.columns:
            if str(c).strip().lower() == kw.lower():
                return c
    # 2) 再按候选词优先级做子串匹配
    for kw in candidates:
        for c in df.columns:
            if kw.lower() in str(c).lower():
                return c
    return None

def fmt_hms(seconds):
    """把秒数格式化为 ±H:MM:SS（保留正负号，用于负延时）"""
    seconds = int(round(seconds))
    sign = "-" if seconds < 0 else ""
    seconds = abs(seconds)
    h = seconds // 3600
    m = (seconds % 3600) // 60
    s = seconds % 60
    return f"{sign}{h}:{m:02d}:{s:02d}"

def fmt_hms_series(delta):
    """向量化把秒数序列格式化为 ±H:MM:SS（几万行比逐行 apply 快一个量级）"""
    out = pd.Series("", index=delta.index)
    valid = delta.notna()
    if not valid.any():
        return out
    sec = delta[valid].round().astype("int64")
    sign = pd.Series(np.where(sec < 0, "-", ""), index=sec.index)
    a = sec.abs()
    strs = (sign + (a // 3600).astype(str) + ":"
            + (a % 3600 // 60).astype(str).str.zfill(2) + ":"
            + (a % 60).astype(str).str.zfill(2))
    out[valid] = strs
    return out

def pick_input_file():
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title="选择 GPS 数据文件",
        filetypes=[("Excel/CSV", "*.xlsx;*.xls;*.csv"), ("All", "*.*")])
    root.destroy()
    return path

def ask_acc_thresholds(default_on=15, default_off=120):
    """弹窗分别输入 ACC 开 / ACC 关 的丢失判定间隔（秒）"""
    import tkinter as tk
    from tkinter import simpledialog
    root = tk.Tk()
    root.withdraw()
    on = simpledialog.askfloat(
        "ACC 开 间隔阈值",
        "请输入【ACC 开】时相邻上报允许的最大间隔（秒）：\n超过该值将标记为「丢失」",
        initialvalue=default_on)
    off = simpledialog.askfloat(
        "ACC 关 间隔阈值",
        "请输入【ACC 关】时相邻上报允许的最大间隔（秒）：\n超过该值将标记为「丢失」",
        initialvalue=default_off)
    root.destroy()
    return on, off

def parse_acc(series):
    """把 ACC 状态列解析为 'on' / 'off'；兼容 "ACC 开 ,左转向灯信号"、"ACC 关"、ON/OFF、1/0 等；
    优先取 "ACC" 后紧跟的开/关，无法识别的值按 on 处理（用更严阈值，宁多报不漏报）"""
    s = series.astype(str).str.strip().str.lower()
    norm = s.str.replace(r"\s+", "", regex=True)          # 压掉空格
    # 优先：acc 后面紧跟的状态词
    m = norm.str.extract(r"acc(开|关|on|off|1|0)", expand=False)
    state = pd.Series("on", index=series.index, dtype=object)
    state[m.isin(["关", "off", "0"])] = "off"
    # 没有 acc 前缀的，退回整体判断
    rest = m.isna()
    if rest.any():
        sub = norm[rest]
        off = sub.str.contains("关|off|false|close", regex=True) | sub.isin(["0"])
        state.loc[sub[off].index] = "off"
    return state

def parse_pos(series):
    """把定位状态列解析为 'pos'（已定位）/ 'nopos'（未定位）；
    兼容 定位/未定位、已定位/未定位、A/V、1/0、true/false 等写法，无法识别的按已定位处理"""
    s = series.astype(str).str.strip().str.lower()
    norm = s.str.replace(r"[^a-z0-9\u4e00-\u9fff]", "", regex=True)
    neg = (norm.str.contains("未|无|不|失效|无效|invalid|nofix|novalid", regex=True)
           | norm.isin(["0", "v", "false", "n", "no"]))
    state = pd.Series("pos", index=series.index, dtype=object)
    state[neg] = "nopos"
    return state

def load_df(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xls"):
        return pd.read_excel(path)
    return pd.read_csv(path)

def style_output(path, min_width=8, max_width=60, sample_rows=300):
    """单次读写完成全部美化（仅 xlsx），几万行数据只 load/save 工作簿一次：
    1) 「延时」列负值标红加粗  2) 「丢失」列标黄加粗
    3) 自适应列宽（中文计 2 宽，采样前 N 行）  4) 冻结首行表头
    """
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from datetime import datetime, date, time as dtime
    wb = load_workbook(path)
    ws = wb.active

    # 定位「延时」「丢失」列
    delay_col = lost_col = None
    for cell in ws[1]:
        if str(cell.value) == "延时":
            delay_col = cell.column
        elif str(cell.value) == "丢失":
            lost_col = cell.column

    # 逐行扫描：负延时标红、丢失标黄（一次循环同时处理两列）
    red_font = Font(color="FF0000", bold=True)
    red_fill = PatternFill(fill_type="solid", start_color="FFFFC7CE")
    bold_font = Font(color="FF9C6500", bold=True)
    yellow_fill = PatternFill(fill_type="solid", start_color="FFFFEB9C")
    n_neg = n_lost = 0
    for row in range(2, ws.max_row + 1):
        if delay_col:
            c = ws.cell(row=row, column=delay_col)
            if isinstance(c.value, str) and c.value.startswith("-"):
                c.font = red_font
                c.fill = red_fill
                n_neg += 1
        if lost_col:
            c = ws.cell(row=row, column=lost_col)
            if c.value == "丢失":
                c.font = bold_font
                c.fill = yellow_fill
                n_lost += 1

    # 自适应列宽（表头 + 前 N 行采样，避免大文件变慢）
    widths = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, sample_rows + 1)):
        for cell in row:
            v = cell.value
            if v is None:
                continue
            s = "0000-00-00 00:00:00" if isinstance(v, (datetime, date, dtime)) else str(v)
            w = sum(2 if ord(ch) > 127 else 1 for ch in s)
            if w > widths.get(cell.column, 0):
                widths[cell.column] = w
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(w + 3, max_width))

    ws.freeze_panes = "A2"
    wb.save(path)
    return n_neg, n_lost

def safe_write(df, out):
    """写出结果文件；若目标被占用（如在 Excel/WPS 中打开）则自动加时间戳改名保存"""
    from datetime import datetime
    try:
        if out.lower().endswith((".xlsx", ".xls")):
            df.to_excel(out, index=False)
        else:
            df.to_csv(out, index=False, encoding="utf-8-sig")
        return out
    except PermissionError:
        base, ext = os.path.splitext(out)
        alt = f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext or '.xlsx'}"
        if ext.lower() == ".xlsx":
            df.to_excel(alt, index=False)
        else:
            df.to_csv(alt, index=False, encoding="utf-8-sig")
        print(f"提示：{out} 正被占用（多半在 Excel/WPS 中打开着），已改存为：{alt}")
        print("      关闭该文件后重跑即可覆盖原文件。")
        return alt

def analyze(df, acc_on_sec=15, acc_off_sec=120, sort_by_time=False):
    sat_col = find_col(df, *SAT_COL_CANDIDATES)
    recv_col = find_col(df, *RECV_COL_CANDIDATES)
    acc_col = find_col(df, *ACC_COL_CANDIDATES)
    if not sat_col or not recv_col:
        raise RuntimeError(
            f"未找到时间列。现有列：{list(df.columns)}\n"
            f"请确认文件含「卫星时间」「接收时间」列，或在脚本顶部修改 SAT_COL_CANDIDATES / RECV_COL_CANDIDATES。")
    if not acc_col:
        raise RuntimeError(
            f"未找到 ACC/车辆状态列。现有列：{list(df.columns)}\n"
            f"请确认文件含「ACC」状态列，或在脚本顶部修改 ACC_COL_CANDIDATES。")

    df[sat_col] = pd.to_datetime(df[sat_col], errors="coerce")
    df[recv_col] = pd.to_datetime(df[recv_col], errors="coerce")

    if sort_by_time:
        df = df.sort_values(sat_col, kind="stable").reset_index(drop=True)

    # 1) 延时：每行都计算 接收 - 卫星，保留正负号（接收早于卫星即为负，向量化计算）
    delta = (df[recv_col] - df[sat_col]).dt.total_seconds()
    df["延时"] = fmt_hms_series(delta)

    # 2) 丢失：按当前行 ACC 状态选用阈值——ACC 关用 acc_off_sec，其余用 acc_on_sec
    gap = df[sat_col].diff().dt.total_seconds()   # 首行为 NaN
    acc_state = parse_acc(df[acc_col])
    limit = pd.Series(
        np.where(acc_state.to_numpy() == "off", float(acc_off_sec), float(acc_on_sec)),
        index=df.index)
    df["丢失"] = ""
    gap_mask = gap.abs() > limit
    df.loc[gap_mask, "丢失"] = "丢失"
    df.loc[~gap_mask & gap.notna(), "丢失"] = "正常"

    # 3) 重发：逐行比较卫星时间，与上一行相同则标记重发
    df["重发"] = ""
    same_as_prev = df[sat_col].notna() & df[sat_col].eq(df[sat_col].shift(1))
    df.loc[same_as_prev, "重发"] = "重发"
    return df

def main():
    t0 = time.time()
    parser = argparse.ArgumentParser(description="GPS 延时/丢失分析（按 ACC 状态分档判丢失）")
    parser.add_argument("--input", help="输入 xlsx/csv 路径")
    parser.add_argument("--acc-on", type=float, help="ACC 开 间隔阈值(秒)，默认 15")
    parser.add_argument("--acc-off", type=float, help="ACC 关 间隔阈值(秒)，默认 120")
    parser.add_argument("--output", help="输出路径(可选)")
    args = parser.parse_args()

    input_path = args.input
    if not input_path:
        try:
            input_path = pick_input_file()
        except Exception:
            input_path = input("请输入数据文件路径：").strip('"')
    if not input_path or not os.path.exists(input_path):
        print("未选择有效文件，退出。")
        sys.exit(1)

    acc_on = args.acc_on
    acc_off = args.acc_off
    if acc_on is None or acc_off is None:
        try:
            t_on, t_off = ask_acc_thresholds(15, 120)
            acc_on = acc_on if acc_on is not None else (t_on if t_on is not None else 15.0)
            acc_off = acc_off if acc_off is not None else (t_off if t_off is not None else 120.0)
        except Exception:
            acc_on = acc_on if acc_on is not None else float(input("ACC 开 间隔阈值(秒)，默认 15：") or 15)
            acc_off = acc_off if acc_off is not None else float(input("ACC 关 间隔阈值(秒)，默认 120：") or 120)

    df = load_df(input_path)
    df = analyze(df, acc_on, acc_off)

    out = args.output
    if not out:
        base, ext = os.path.splitext(input_path)
        out = base + "_分析.xlsx"
    out = safe_write(df, out)
    n_neg = n_lost_marked = 0
    if out.lower().endswith(".xlsx"):
        n_neg, n_lost_marked = style_output(out)
    # CSV 无法保存标色样式

    total = len(df)
    n_neg = int(df["延时"].astype(str).str.startswith("-").sum())
    n_lost = int((df["丢失"] == "丢失").sum())
    n_resend = int((df["重发"] == "重发").sum())
    acc_state = parse_acc(df[find_col(df, *ACC_COL_CANDIDATES)])
    n_lost_on = int(((df["丢失"] == "丢失") & (acc_state == "on")).sum())
    n_lost_off = int(((df["丢失"] == "丢失") & (acc_state == "off")).sum())

    def pct(n):
        return f"占总数据比例 {n / total * 100:.2f}%" if total else "占总数据比例 0.00%"

    print(f"处理完成：共 {total} 行，耗时 {time.time() - t0:.1f} 秒")
    print(f"  总行数：{total}")
    print(f"  延时为负(标红) 行数：{n_neg}   {pct(n_neg)}")
    print(f"  丢失(标黄) 行数：{n_lost}   {pct(n_lost)}"
          f"   [ACC开>{acc_on:g}s: {n_lost_on} 行 | ACC关>{acc_off:g}s: {n_lost_off} 行]")
    print(f"  重发 行数：{n_resend}   {pct(n_resend)}")

    # 定位状态统计（列不存在则跳过，不影响分析）
    pos_col = find_col(df, *POS_COL_CANDIDATES)
    if pos_col:
        pos_state = parse_pos(df[pos_col])
        n_pos = int((pos_state == "pos").sum())
        n_nopos = int((pos_state == "nopos").sum())
        print(f"  定位   行数：{n_pos}   {pct(n_pos)}")
        print(f"  未定位 行数：{n_nopos}   {pct(n_nopos)}")
    else:
        print("  定位统计：未找到定位状态列，已跳过"
              "（可在脚本顶部 POS_COL_CANDIDATES 中补充列名）")
    print(f"  输出文件：{out}")

if __name__ == "__main__":
    main()
